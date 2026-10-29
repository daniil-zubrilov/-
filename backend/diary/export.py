import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from .models import Workout

class ExportExcelView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        # Создаем книгу Excel
        wb = Workbook()
        ws = wb.active
        ws.title = "Мои тренировки"

        # Заголовки
        headers = ["Дата", "Название тренировки", "Заметки", "Упражнение", "Вес (кг)", "Повторы"]
        ws.append(headers)

        # Стилизация заголовков
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="FF5A1F", end_color="FF5A1F", fill_type="solid")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # Берем только тренировки текущего пользователя
        workouts = Workout.objects.filter(owner=request.user).prefetch_related('exercises__sets')

        for w in workouts:
            if not w.exercises.exists():
                ws.append([w.date, w.title, w.notes, "", "", ""])
            
            for ex in w.exercises.all():
                if not ex.sets.exists():
                    ws.append([w.date, w.title, w.notes, ex.name, "", ""])
                for s in ex.sets.all():
                    ws.append([w.date, w.title, w.notes, ex.name, s.weight, s.reps])

        # Настройка ширины колонок
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 30

        # Сохраняем в память
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        # Отдаем файл
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="fitness_export.xlsx"'
        return response